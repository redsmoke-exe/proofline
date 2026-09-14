from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.pdfgen.canvas import Canvas

from cv_agent.ingestion import UniversalDocumentReader, ingest_request, normalize_text
from cv_agent.schemas import IngestionRequest


def test_normalize_text_repairs_common_extraction_noise() -> None:
    assert normalize_text("alpha-\nbeta\r\n\r\n\r\ngamma\x00") == "alphabeta\n\ngamma"


def test_ingest_raw_jd_and_deduplicate_candidate_files(tmp_path: Path) -> None:
    cv = tmp_path / "cv.md"
    cv.write_text("# Person\n\nExperience with Python.", encoding="utf-8")
    bundle = ingest_request(
        IngestionRequest(
            jd_path=None,
            jd_text="Python engineer needed",
            cv_paths=[cv, cv],
            max_chars_per_document=10_000,
        )
    )
    assert bundle.job_description.text == "Python engineer needed"
    assert len(bundle.candidate_documents) == 1
    assert bundle.candidate_documents[0].kind == "md"


def test_docx_and_pdf_are_extracted(tmp_path: Path) -> None:
    docx_path = tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("Candidate Name")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "SQL"
    document.save(docx_path)

    pdf_path = tmp_path / "job.pdf"
    canvas = Canvas(str(pdf_path))
    canvas.drawString(72, 720, "Data Platform Engineer with Airflow")
    canvas.save()

    reader = UniversalDocumentReader()
    assert "Python | SQL" in reader.read(docx_path).text
    assert "Data Platform Engineer" in reader.read(pdf_path).text
